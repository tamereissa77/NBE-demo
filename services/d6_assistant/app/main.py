"""D6 · AI Assistant (PLAN.md §4 D6).

One assistant for file chat, drafting, summarisation, translation and code explanation. A router picks
the model per task and says why; the portal shows that as the Model-Selection Panel.

Two front doors:
  /chat, /upload, /route      portal (behind the gateway at /api/d6)
  /v1/chat/completions, /v1/models   OpenAI-compatible, model "nbe-auto", for Open WebUI and Dataiku LLM Mesh
                                      (gateway path /v1; bearer key D6_API_KEY when set)
"""
import asyncio
import base64
import io
import json
import os
import re
import time
import uuid
from pathlib import Path

import httpx
from fastapi import File, Header, HTTPException, Request, UploadFile
from fastapi.responses import Response, StreamingResponse
from PIL import Image, ImageOps
from pydantic import BaseModel

from app import coding, conversation, drafting, evidence, intake, sandbox, sheets, summarising, translating
from nbe_core import create_app, llm, platform, retrieval, settings
from nbe_core.audit import log_event
from nbe_core.text import detect_lang
from nbe_core.trace import Trace

DOCS_DIR = Path("/data/generated/d6")
COLLECTION = "d6_docs"
API_KEY = os.getenv("D6_API_KEY", "")
app = create_app("d6", f"{settings.bank_short} AI Assistant", phase="wave1")
FILES: dict[str, dict] = {}       # uploaded files for file chat (in memory, demo scope)

# Reading uploads. Pages with no usable text layer go to the vision model: scans and photos have none,
# and Arabic text layers come out in visual order (reversed digits, scrambled table cells).
IMAGE_TYPES = (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff")
MIN_PAGE_CHARS = 40               # fewer letters and digits than this in a page's text layer: the page is a scan
MAX_VISION_PAGES = 15             # demo budget: each page is one vision call
MAX_IMAGE_SIDE = 2200             # phone photos are downscaled to this before they are sent
TRANSCRIBE_TOKENS = 3000
TRANSCRIBE = ("Transcribe all the text on this page exactly as written, in its original language or languages. "
              "Do not translate, summarise, correct or comment. Keep the reading order. Put # before a section title only: a short line that names "
              "the section under it. Dates, greetings, signatures and numbered paragraphs are not titles. "
              "Write tables as Markdown tables. "
              "Keep every number, date and reference exactly as it appears. If the page has no text, reply with nothing.")

# Task -> model alias, and why. "hero" falls back to the workhorse at the gateway when it is not running.
ROUTES = {
    "file_chat": ("workhorse", "File chat and extraction: long context, Arabic, fast"),
    "drafting": ("workhorse", "Drafting emails and letters: fluent Arabic and English"),
    "summarise": ("workhorse", "Summaries and key points: fast, good Arabic"),
    "translation": ("workhorse", "Translation Arabic ↔ English: strongest Arabic among fast models"),
    "code": ("hero", "Code: gpt-oss-120b reasons best about code; answers grounded in local docs, checked and run before they are shown"),
    "general": ("workhorse", "General questions: the fast default"),
}
RULES = [  # cheap first pass; the LLM classifier handles anything the rules do not catch
    ("code", r"```|\bdef |\bclass |\bimport |\bSELECT\b|\bpandas\b|\bnumpy\b|\bpython\b|\bSQL\b|\bdf\.|\.groupby\(|\bVBA\b|\bregex\b|\bexcel formula\b|كود|برمج|دالة|الكود|معادلة"),
    ("translation", r"\btranslat|ترجم|ترجمة|للإنجليزي|للانجليزي|للعربي|بالإنجليزية|into english|into arabic"),
    ("summarise", r"\bsummar|لخص|لخّص|تلخيص|key points|النقاط الرئيسية|أهم النقاط"),
    ("drafting", r"\bdraft|\b(write|prepare|compose) (an?|the) (short |brief |formal |official |internal )*(e-?mail|letter|memo|reply|response|notice|circular|invitation|announcement)"
                 r"|اكتب(لي)? (إيميل|ايميل|بريد|خطاب|رسالة|رد|مذكرة|تعميم|إشعار|دعوة)|حض[ّ]?ر (مذكرة|رد|خطاب|بريد|دعوة)|صياغة|صيغ|رد(اً|ا)? رسمي"),
]


def require_key(authorization: str | None) -> None:
    if API_KEY and authorization != f"Bearer {API_KEY}":
        raise HTTPException(401, "invalid API key")


TASK_HELP = ("drafting: write or change an email, letter, memo or message (shorter, longer, another tone, add a line); "
             "summarise: condense a text the employee pasted; translation: between Arabic and English; "
             "code: programming, SQL, or questions about code and how it behaves; general: anything else")


# An instruction that changes the draft on the table ("make it shorter", "add a line about...", "غيّر الخاتمة")
REVISE = re.compile(r"^\W*(please\s+)?(change|make|add|remove|delete|replace|shorten|lengthen|rewrite|rephrase|reword|fix|correct|use|keep|put|move|sign)\b"
                    r"|^\W*(غيّر|غير|اجعل|أضف|اضف|احذف|استبدل|اختصر|أعد|اعد|عدّل|عدل|صحّح|صحح|اكتبه|اجعله)", re.I)


# Code asked for outright. With a file attached, only this makes a request a code request ("what is the branch code?" is a question about the file)
CODE_ASK = re.compile(r"```|\b(pandas|python|numpy|sql|vba|regex|dataframe)\b|\b(write|generate|give me|create|need|want)\b.{0,40}\b(code|script|query|formula|function|macro)\b"
                      r"|(اكتب|اعمل|أعطني|اعطني|أريد|اريد|عايز).{0,40}(كود|سكربت|استعلام|معادلة|دالة|برنامج)|بايثون", re.I | re.S)
CONTINUES = re.compile(r"^\W*(now|and|also|then|same|again|what about|how about)\b|^\W*(والآن|الآن|ثم|نفس|كذلك|أيضا|أيضاً|وماذا عن|ماذا عن)", re.I)


def rule_task(text: str) -> str | None:
    """The task whose keyword comes first in the message. "Draft an e-mail that passes this summary on" is drafting,
    "summarise this draft" is a summary: the word that opens the request names the task, later ones describe its subject."""
    hits = [(m.start(), i, name) for i, (name, pattern) in enumerate(RULES) if (m := re.search(pattern, text, re.I))]
    return min(hits)[2] if hits else None


async def route(text: str, has_file: bool, prev: dict | None = None) -> dict:
    """Language + task + model, with the reason. Rules first, then a small LLM classification.
    `prev` ({question, task, lang}) is the previous request of the conversation: a follow-up such as "make it
    shorter" or "why is that faster?" names no task of its own and belongs to the request it follows."""
    lang = conversation.language(text, prev and prev.get("lang"))
    # "draft a reply to this letter", "summarise this file": the task named, with the file as its source
    own_task = rule_task(text) in ("drafting", "summarise") or bool(prev and prev.get("task") in ("drafting", "summarise") and REVISE.search(text))
    # code asked for with a file attached, or the code request before this one carried on ("now do it per product")
    wants_code = rule_task(text) in (None, "code") and bool(CODE_ASK.search(text) or (
        has_file and prev and prev.get("task") == "code" and not rule_task(text) and (REVISE.search(text) or CONTINUES.search(text))))
    if has_file and wants_code:
        task, how = "code", "rule: code asked for, on the attached file"
    elif has_file and not own_task and not re.search(RULES[1][1], text, re.I):
        task, how = "file_chat", "rule: a file is attached"
    else:
        task, how = rule_task(text), ""
        if task:
            how = f"rule: matched {task} keywords"
        elif prev and prev.get("task") in ("drafting", "summarise") and REVISE.search(text):
            task, how = prev["task"], "rule: revises the current " + ("draft" if prev["task"] == "drafting" else "summary")
        else:
            options = [t for t in ROUTES if t != "file_chat"]          # nothing is attached, so it cannot be file chat
            carried = prev["task"] if prev and prev.get("task") in options else None
            before = (f"The previous request in this conversation was a {carried} request: \"{prev['question'][:400]}\"\n"
                      f"If the latest message continues it, changes its result or asks about it or its answer, the task is still {carried}. "
                      "If it starts something new, classify it on its own.\n") if carried else ""
            try:
                data, _ = await llm.chat_json([{"role": "user", "content":
                    f"Classify the employee's latest message into one task.\n{TASK_HELP}\n{before}"
                    f"Answer JSON only.\nLatest message: {text[:1500]}"}],
                    {"type": "object", "properties": {"task": {"type": "string", "enum": options}}, "required": ["task"]}, max_tokens=20)
                task = data["task"]
                how = f"LLM classifier: continues the previous {task} request" if task == carried else "LLM classifier (workhorse)"
            except (httpx.HTTPError, llm.JSONOutputError, KeyError):
                # The classifier is a convenience: a failed call must not fail the chat
                task, how = "general", "classifier unavailable: general by default"
    alias, reason = ROUTES[task]
    return {"lang": lang, "task": task, "alias": alias, "reason": reason, "how": how}


async def code_context(question: str, trace: Trace, previous: str = "") -> str:
    """`previous` is the question this one follows up on: "and what does the endpoint argument do?" names no
    function, so the documentation is searched with the earlier question as well."""
    query = f"{previous[:600]}\n{question}" if previous else question
    with trace.step("retrieval", "Ground code answer in local pandas/numpy docs" + (" (searched with the previous question)" if previous else "")) as s:
        try:
            hits = await retrieval.search(COLLECTION, query, k=3, candidates=12)
        except httpx.HTTPError:
            hits = []
        s["detail"] = [h.get("name") for h in hits]
    if not hits:
        return ""
    return "Reference documentation (use only APIs documented here or in the standard library):\n" + \
        "\n\n".join(f"### {h['name']}\n{h['text'][:1800]}" for h in hits)


SYSTEM = {
    "ar": (f"أنت مساعد موظفي {settings.bank_ar}. أجب بالعربية الفصحى الواضحة ما لم يُطلب غير ذلك. كن دقيقاً ومختصراً، "
           "ولا تخترع معلومات عن البنك. عند شرح الأكواد اكتب المصطلحات التقنية بالإنجليزية."),
    "en": f"You are {settings.bank_en}'s internal employee assistant. Be accurate and concise; never invent bank facts.",
}


# Per-task output rules. The Arabic gate showed the models add preambles and several alternatives to drafts.
TASK_RULES = {
    "drafting": "",          # the drafting rules, calendar and bank facts are built per request: app/drafting.py
    "translation": "Output only the translation, with no quotes, labels or commentary.",
    "summarise": "",         # the summary rules and the employee's choices are built per request: app/summarising.py
    "code": "Explain clearly, point out real bugs, and only use APIs that exist in the documented library versions.",
}


def build_messages(messages: list[dict], r: dict, extra_context: str = "", draft: dict | None = None) -> list[dict]:
    """`draft` ({opts, current}) carries the employee's choices for a drafting request and, for a revision, the draft being revised."""
    draft = draft or {}
    if r["task"] == "summarise":
        sys = SYSTEM[r["lang"]].split(" عند شرح الأكواد")[0] + " " + drafting.date_line(r["lang"]) + "\n\n" + summarising.context(r["lang"], draft.get("opts")) + \
            ("\n\n" + extra_context if extra_context else "") + ("\n\n" + summarising.reminder(r["lang"], draft.get("opts")) if summarising.reminder(r["lang"], draft.get("opts")) else "")
        return [{"role": "system", "content": sys}, *[m for m in messages if m["role"] != "system"]]
    task = "\n\n" + drafting.context(r["lang"], draft.get("opts"), draft.get("current", "")) if r["task"] == "drafting" else \
        "\n\n" + draft["code"] if r["task"] == "code" and draft.get("code") else " " + TASK_RULES[r["task"]] if r["task"] in TASK_RULES else ""
    sys = SYSTEM[r["lang"]] + " " + drafting.date_line(r["lang"]) + task + ("\n\n" + extra_context if extra_context else "")
    return [{"role": "system", "content": sys}, *[m for m in messages if m["role"] != "system"]]


FILE_CHARS = 24000                # of one file shown to the model
FILES_CHARS = int(os.getenv("D6_FILES_CHARS", "60000"))     # of all attached files together
SHEET_NOTE = ("Each sheet starts with figures computed exactly from all its rows: for a total, average, count or total per "
              "category, quote those figures and do not add up the rows yourself.")


# Trust and evidence (app/evidence.py): numbered passages, citations under each answer, checks on figures and quotes.
# D6_EVIDENCE=0 switches all of it off: the files then go to the model exactly as they did before.
EVIDENCE = os.getenv("D6_EVIDENCE", "1") == "1"
CITE = ("The text is divided into numbered passages: [S1], [S2], ... Every sentence that states something from a file must "
        "end with the number of the passage it comes from, in square brackets exactly like this: [S4]. Cite only a passage "
        "that states the fact, and never a number that is not in the text. Write the answer in your own words, in the language of the question, and keep "
        "names, codes and figures exactly as written; use quotation marks only when the user asks for the exact wording.")


def files_context(files: list[dict], trace: Trace, cite: bool = False) -> tuple[str, dict]:
    """The attached files as the model sees them, with the instruction on how to answer from them.
    With `cite`, the text carries passage numbers; the second value is those passages by number."""
    for f in files:
        info = f.get("info") or {}
        trace.add("file", f"{f['name']}: read by {info.get('converted_by', '?')}",
                  {"chars": len(f["text"]), "passages": len(f.get("passages") or []), "structure": info.get("structure"),
                   "read_in_s": info.get("read_s")}, int((info.get("read_s") or 0) * 1000))
    by_label: dict = {}
    if len(files) == 1:
        f = files[0]
        shown = f["text"][:FILE_CHARS]
        if cite:
            [shown], by_label = evidence.number(files, [FILE_CHARS])
            trace.add("files", f"1 file in the prompt as {len(by_label)} numbered passages",
                      [{"file": f["name"], "passages_shown": len(by_label), "passages_total": len(f.get("passages") or [])}])
        where = (f"The file is a spreadsheet, divided by [Sheet: name] markers. {SHEET_NOTE} Answer from the file when the "
                 "question is about it, and name the sheet the answer is in." if f.get("kind") == "sheets" else
                 "The file is divided by [Page N] markers. Answer from the file when the question is about it, and name the page "
                 "(and the heading, if there is one) where the answer is." if f.get("paged") else
                 "Answer from the file when the question is about it, and say where in the file the answer is.")
        return f"Attached file: {f['name']}\n---\n{shown}\n---\n{where}" + (f"\n{CITE}" if cite else ""), by_label
    # Several files share one budget. Short files keep all their text; what they do not use goes to the longer ones.
    share, left, room = {}, sorted(range(len(files)), key=lambda i: len(files[i]["text"])), FILES_CHARS
    for n, i in enumerate(left):
        share[i] = min(len(files[i]["text"]), FILE_CHARS, room // (len(left) - n))
        room -= share[i]
    # How each file was read goes in its header, so "the scanned letter" or "the spreadsheet" can be told apart
    how = lambda f: f" ({f['info']['converted_by']})" if f.get("info") else ""  # noqa: E731
    shown = [f["text"][:share[i]] for i, f in enumerate(files)]
    if cite:
        shown, by_label = evidence.number(files, [share[i] for i in range(len(files))])
    blocks = [f"=== File {i + 1}: {f['name']}{how(f)} ===\n{shown[i]}" for i, f in enumerate(files)]
    trace.add("files", f"{len(files)} files in the prompt" + (f" as {len(by_label)} numbered passages" if cite else ""),
              [{"file": f["name"], "chars_shown": share[i], "chars_total": len(f["text"])} for i, f in enumerate(files)])
    rules = (f"{len(files)} files are attached, each between \"=== File N: name ===\" markers; inside a file, [Page N] or "
             "[Sheet: name] markers show where you are. Answer from the files when the question is about them.\n"
             "- For every fact, name the file it comes from, and the page or sheet when there is one.\n"
             "- When a question spans files, take each part from the file that has it.\n"
             "- When files give different values for the same thing, do not pick one: give each value with its file.\n"
             "- When none of the files has the answer, say so.")
    if any(f.get("kind") == "sheets" for f in files):
        rules += f"\n- In a spreadsheet: {SHEET_NOTE}"
    if cite:
        rules += f"\n- {CITE} The numbers run on from one file to the next."
    return "Attached files:\n\n" + "\n\n".join(blocks) + "\n---\n" + rules, by_label


class ChatIn(BaseModel):
    messages: list[dict]
    file_id: str | None = None
    file_ids: list[str] = []                # several files attached together; `file_id` is the single-file form
    evidence: bool | None = None            # numbered passages and checked citations; default: the D6_EVIDENCE setting
    draft_options: dict = {}                # for a drafting request: type, tone, length, audience, language, template
    summary_options: dict = {}              # for a summary: form, length, audience, language, focus
    translation_options: dict = {}          # for a translation: target, register, numerals
    code_options: dict = {}                 # for code: language or dialect, depth of the explanation, its language
    conversation_id: str | None = None      # with it, the server keeps the history and `messages` is just the new question


@app.post("/conversation")
def new_conversation():
    return {"conversation_id": conversation.new()}


@app.get("/conversation/{cid}")
def get_conversation(cid: str):
    conv = conversation.get(cid)
    if not conv:
        raise HTTPException(404, {"code": "conversation_gone", "message": "This conversation is no longer available."})
    return conversation.view(conv)


@app.post("/chat")
async def chat(body: ChatIn):
    trace = Trace()
    # Passwords, keys and connection-string credentials are taken out before any model sees the message (app/coding.py)
    question, hidden = coding.mask(body.messages[-1]["content"])
    hidden = [None] * coding.places(question) if hidden else []          # counted by place: a password and the connection string that repeats it are two
    if hidden:
        trace.add("code", f"{len(hidden)} secret(s) replaced by placeholders before the message was sent to a model")
        body.messages = [{**m, "content": coding.mask(m["content"])[0]} if isinstance(m.get("content"), str) else m for m in body.messages]
    ids = list(dict.fromkeys(body.file_ids or ([body.file_id] if body.file_id else [])))
    if len(ids) > intake.MAX_FILES:
        no = intake.Rejected("too_many_files", conversation.language(question, None), limit=intake.MAX_FILES)
        raise HTTPException(no.status, no.detail)
    files = [FILES[i] for i in ids if i in FILES]
    conv = conversation.get(body.conversation_id)
    if body.conversation_id and not conv:
        raise HTTPException(404, {"code": "conversation_gone", "message": "This conversation is no longer available."})
    if conv:
        prev = conv["last"]
    else:       # a caller that sends its own history: the previous request is the user message before this one
        earlier = [m["content"] for m in body.messages[:-1] if m.get("role") == "user" and isinstance(m.get("content"), str)]
        prev = {"question": earlier[-1], "task": rule_task(earlier[-1]), "lang": conversation.language(earlier[-1], None)} if earlier else None
    with trace.step("route", "Detect language and task, choose model") as s:
        r = await route(question, bool(files), prev)
        s["detail"] = r
    if r["task"] == "translation":
        return await translate(body, question, files, conv, r, trace)
    history, summary = await conversation.context(conv, trace) if conv else (body.messages[:-1], "")
    extra, code_rules = "", ""
    is_draft, is_summary, is_code = r["task"] == "drafting", r["task"] == "summarise", r["task"] == "code"
    opts = {k: v for k, v in (body.draft_options if is_draft else body.summary_options if is_summary else body.code_options if is_code else {}).items() if isinstance(v, str)}
    if (is_draft or is_summary) and opts.get("language") in ("ar", "en"):
        r["lang"] = opts["language"]
    data_files = [f for f in files if f.get("source")] if is_code else []       # spreadsheets the code can be run on
    if is_code:
        if opts.get("explain") in ("ar", "en"):
            r["lang"] = opts["explain"]
        follows = bool(prev) and prev.get("task") == "code" and (r["how"].startswith("LLM classifier: continues") or len(question.split()) < 12)
        extra = await code_context(question, trace, prev["question"] if follows else "")
        missing = coding.unknown_references(question)
        code_rules = coding.context(r["lang"], opts, coding.schema(data_files) if data_files else "", coding.warnings(question, "", len(hidden), r["lang"]), missing)
        if missing:
            trace.add("code", "Names in the request checked against the installed libraries: " + ", ".join(missing) + " not found", missing)
    ask = question                              # what the model is asked; for a summary of pasted text, the request without the text
    if is_summary:
        if not files:
            ask, text_source = summarising.pasted(question, r["lang"])
            if text_source:
                if opts.get("language") not in ("ar", "en"):     # the language of the request, not of the text pasted under it
                    r["lang"] = conversation.language(ask, prev and prev.get("lang"))
                    text_source["name"] = "النص الملصق" if r["lang"] == "ar" else "pasted text"
                files = [text_source]
            elif conv and conv.get("summary_source") and r["how"].startswith(("rule: revises", "LLM classifier: continues")):
                files = conv["summary_source"]              # "make it shorter": the same text again
            elif not history and len(question.split()) < 25:     # nothing pasted, nothing attached, nothing said before
                r["model"], r["fallback"] = r["alias"], False
                return {"answer": summarising.NOTHING[r["lang"]], "route": r, "meta": {"latency_ms": 0, "tokens_per_s": 0, "completion_tokens": 0},
                        "trace": trace.to_list(), "conversation_id": body.conversation_id}
        if conv is not None and files:
            conv["summary_source"] = files
    passages: dict = {}
    in_sections = 0
    if files and is_summary and (EVIDENCE if body.evidence is None else body.evidence):
        extra, passages, in_sections = await summary_context(files, opts, r["lang"], trace)
    elif files and is_code:
        # A spreadsheet is described by its sheets, columns and values, not sent row by row: the code is run on the file itself
        other = [f for f in files if not f.get("source")]
        if other:
            shown, _ = files_context(other, trace, False)
            extra = shown + (f"\n\n{extra}" if extra else "")
        if data_files:
            trace.add("files", f"{len(data_files)} spreadsheet(s) described to the model by columns and values; the rows stay here",
                      [{"file": f["name"], "sheets": [{"name": s["name"], "rows": s["rows"], "header_row": s["header_row"]} for s in f["sheets"]]} for f in data_files])
    elif files:
        # A draft is not asked to carry passage marks (a letter full of [S4] is not a letter): its sources are found afterwards
        extra, passages = files_context(files, trace, (EVIDENCE if body.evidence is None else body.evidence) and not is_draft)
    if summary:
        extra = f"Notes on the earlier part of this conversation (those messages are no longer shown in full):\n{summary}" + \
            (f"\n\n{extra}" if extra else "")
    # A follow-up that continues a drafting request revises the latest version (which the employee may have edited)
    versions = conv["drafts"] if conv else []
    revising = is_draft and bool(versions) and r["how"].startswith(("LLM classifier: continues", "rule: revises"))
    kwargs = {"reasoning_effort": "low"} if r["alias"] == "hero" else {}
    messages = [*history, {"role": "user", "content": ask}]
    system = build_messages(messages, r, extra, {"opts": opts, "current": versions[-1]["text"] if revising else "", "code": code_rules})
    res = await llm.chat(system, r["alias"], max_tokens=2500 if is_code else 1500, temperature=0.2, **kwargs)
    trace.llm(f"Answer ({r['alias']})", res)
    if is_code and (not res.text.strip() or res.text.count("```") % 2):
        # An answer with no text, or one that stops inside a code block, is not shown. It happens when the code model is
        # away and its stand-in spends the whole allowance thinking: the request goes once more, to the workhorse, plainly.
        res = await llm.chat(system, "workhorse", max_tokens=3000, temperature=0.2)
        trace.llm("Answer again (workhorse): the first came back empty or cut off inside the code", res)
    r["model"], r["fallback"] = res.model, res.model != r["alias"]
    answer, proof = res.text, {}
    if is_summary and passages:
        answer, report = await finish_summary(answer, passages, files, opts, r["lang"], in_sections, trace)
        proof["summary"] = report
    if passages:
        answer = evidence.tidy(answer, passages)
        sources, quoted = evidence.verify(answer, passages), evidence.quotes(answer, passages)
        if is_summary and not sources:              # a table of action items rarely carries marks: its sources are found by its figures
            sources = evidence.attribute(answer, passages)
        proof.update(sources=sources, quotes=quoted, evidence=evidence.summary(sources, quoted, evidence.loose_figures(answer, passages)))
        e = proof["evidence"]
        trace.add("evidence", f"{e['cited']} passage(s) cited: {e['verified']} confirmed, {e['to_check']} to check, {e['unchecked']} not checkable"
                  + (f"; {e['quotes_not_found']} of {e['quotes']} quote(s) not found" if e["quotes"] else "")
                  + (f"; figures not in the files: {', '.join(e['figures_not_in_files'])}" if e["figures_not_in_files"] else ""),
                  [{k: s[k] for k in ("label", "file", "where", "status", "missing")} for s in sources] + [{"quote": q["quote"][:80], "found": q["found"]} for q in quoted])
        if not is_summary:
            # A long file is shown to the model only up to its budget. That used to pass unremarked; now the answer says so.
            cover = summarising.coverage(passages, files, 0)
            if not cover["complete"]:
                proof["coverage"] = cover
                trace.add("files", f"Only part of the file(s) was read for this answer: {cover['passages_read']} of {cover['passages_total']} passages", cover)
    if is_code:
        answer, proof["code"] = await finish_code(answer, ask, r, opts, data_files, system[0]["content"], len(hidden), trace)
    if is_draft:
        answer, proof = await finish_draft(answer, question, r, opts, files, conv, body, revising, trace)
    if conv:
        conversation.add(conv, question, answer, {"route": r, "meta": res.meta(), **proof},
                         {"question": ask, "task": r["task"], "lang": r["lang"]}, [f["info"] for f in files if (f.get("info") or {}).get("file_id")])
    return {"answer": answer, "route": r, "meta": res.meta(), "trace": trace.to_list(), "conversation_id": body.conversation_id, **proof}


# ------------------------------------------------------------------ code: checked, and run on the attached data, before it is shown
async def finish_code(answer: str, question: str, r: dict, opts: dict, data_files: list[dict], rules: str, hidden: int, trace: Trace) -> tuple[str, dict]:
    """The model's code answer -> (the answer with code that passed its check, and what it printed; the report under it).
    With a spreadsheet attached the check is the real one: the code is run on a copy of the file (app/sandbox.py).
    Without one it is static (app/coding.py). Code that fails goes back to the model once, with the problem."""
    language = opts.get("language", "")
    wrote = not coding.has_code(question)                 # written for the employee, as opposed to their own code explained
    answer = coding.excel_separators(answer, language)
    data = {f["name"]: f["source"] for f in data_files}

    def target(text: str) -> dict | None:
        found = [b for b in coding.blocks(text) if b["lang"] != "text"]
        python = [b for b in found if b["lang"] == "python"]
        return (python[-1] if python else None) if data else (found[-1] if found and wrote else None)

    async def check(b: dict) -> tuple[str, dict | None]:
        if data:
            with trace.step("code", f"Run the code on {', '.join(data)} (separate process, no network)") as s:
                run = await asyncio.to_thread(sandbox.run, b["code"], data)
                s["detail"] = {"ok": run["ok"], "seconds": run["seconds"], "error": run["error"], "output": run["output"][:400]}
            return ("" if run["ok"] else f"it failed when it was run on the file: {run['error']}"), run
        found = coding.problems(b, language)
        trace.add("code", f"Check the {b['lang']} code before showing it: " + ("; ".join(found) if found else "no problem found"), found)
        return "; ".join(found), None

    b, problem, run, first = target(answer), "", None, ""
    if b:
        problem, run = await check(b)
        if problem:
            first = problem
            try:
                ask = [{"role": "system", "content": rules}, {"role": "user", "content":
                       coding.REPAIR.format(request=question[:1500], lang=b["lang"], code=b["code"], problem=problem)}]
                res = await llm.chat(ask, r["alias"], max_tokens=1500, temperature=0.1, **({"reasoning_effort": "low"} if r["alias"] == "hero" else {}))
                if not coding.blocks(res.text):
                    res = await llm.chat(ask, "workhorse", max_tokens=2000, temperature=0.1)
                trace.llm("Correct the code once, given the problem", res)
                fixed = [x for x in coding.blocks(res.text) if x["lang"] != "text"]
            except httpx.HTTPError:
                fixed = []
            if fixed:
                answer = coding.excel_separators(answer[:b["start"]] + fixed[0]["code"] + "\n" + answer[b["end"]:], language)
                b = target(answer)
                problem, run = await check(b)
    report = {"mode": "write" if wrote else "review", "masked": hidden, "repaired": bool(first) and not problem, "first_problem": first,
              "blocks": [{"lang": x["lang"], "code": x["code"][:4000], "run": run if b and x["start"] == b["start"] else None,
                          # the employee's own faulty code may be quoted back in an explanation: only written code is marked
                          "problems": ([problem] if problem and not run else []) if b and x["start"] == b["start"] else (coding.problems(x, language) if wrote else [])}
                         for x in coding.blocks(answer) if x["lang"] != "text"]}
    if run:
        answer += coding.result_block(run, r["lang"])
    report["warnings"] = coding.warnings(question, answer, hidden, r["lang"])
    if report["warnings"]:
        trace.add("code", "Risks raised by rule: " + ", ".join(w["kind"] for w in report["warnings"]), report["warnings"])
    return answer, report


@app.get("/code/options")
def code_options(lang: str = "en"):
    return coding.options(lang)


# ------------------------------------------------------------------ translation: segment by segment, each one read back
_translate_slots = asyncio.Semaphore(2)
RETRIES = 6                 # segments translated again on their own when a marker was lost or a figure did not come through


async def translate(body: ChatIn, question: str, files: list[dict], conv: dict | None, r: dict, trace: Trace) -> dict:
    """The whole of a pasted text, an attached file or the previous answer, translated in segments and checked (app/translating.py)."""
    opts = {k: v for k, v in body.translation_options.items() if isinstance(v, str)}
    request, pasted_text = translating.request_and_text(question)
    earlier = [m["content"] for m in (conv["messages"] if conv else body.messages[:-1]) if m.get("role") == "assistant" and isinstance(m.get("content"), str)]
    if pasted_text:
        source, what = pasted_text, "pasted text"
    elif earlier and not (files and translating.wants_file(question)):
        # "translate it", "translate into English" after an answer: the answer just given, even with a file still attached.
        # Naming the file ("translate this file") is what selects the file.
        source, what, request = drafting.clean(earlier[-1]), "the previous answer", question
    elif files:
        # A file's own text, part by part; several files one after another, each under its name
        source = "\n\n".join((f"=== {f['name']} ===\n\n" if len(files) > 1 else "") + "\n\n".join(p["text"] for p in f.get("parts") or [{"text": f["text"]}]) for f in files)
        what = ", ".join(f["name"] for f in files)
    else:
        r.update(model=r["alias"], fallback=False)
        return {"answer": translating.NOTHING[r["lang"]], "route": r, "meta": {"latency_ms": 0, "tokens_per_s": 0, "completion_tokens": 0},
                "trace": trace.to_list(), "conversation_id": body.conversation_id}
    to = translating.target(request, source, opts.get("target", ""))
    r["lang"] = to
    language = "English" if to == "en" else "Arabic (Modern Standard Arabic)"
    segs = translating.segments(source)
    for s in segs:
        if s.get("copy"):
            s["target"] = s["text"]
    used, todo = 0, []
    for i, s in enumerate(segs):                # a text beyond the allowance: translate up to it, and say so
        if not s.get("copy") and used + len(s["text"]) <= translating.MAX_CHARS:
            todo.append(i)
            used += len(s["text"])
    stats = {"ms": 0, "tokens": 0, "model": r["alias"]}

    async def ask(content: str, max_tokens: int) -> str:
        async with _translate_slots:
            res = await llm.chat([{"role": "user", "content": content}], r["alias"], max_tokens=max_tokens, temperature=0)
        stats.update(ms=stats["ms"] + res.latency_ms, tokens=stats["tokens"] + res.completion_tokens, model=res.model)
        return res.text

    async def one_batch(idx: list[int]) -> None:
        reply = await ask(translating.prompt([(i + 1, segs[i]["text"]) for i in idx], to, opts), 3500)
        for n, text in translating.parse(reply, [i + 1 for i in idx]).items():
            segs[n - 1]["target"] = text

    groups = [[i for i in g if i in todo] for g in translating.batches(segs)]
    t = time.time()
    try:
        await asyncio.gather(*(one_batch(g) for g in groups if g))
        trace.add("llm", f"Translated {what} into {language}: {len(todo)} segment(s) in {len([g for g in groups if g])} call(s)",
                  {"model": stats["model"], "characters": used, "completion_tokens": stats["tokens"]}, int((time.time() - t) * 1000))
        # Read every segment back; one that lost its marker or a figure is translated again on its own, with the figures named
        retried = 0
        for i in todo:
            s = segs[i]
            s["check"] = translating.check(s["text"], s.get("target", ""), to)
            if s["check"]["status"] != "ok" and retried < RETRIES:
                retried += 1
                must = translating._figures(s["text"])
                terms = translating.terms_for(s["text"], to)
                again = await ask(f"Translate this into {language}. Output only the translation, in the same layout."
                                  + (f" These must appear exactly as written: {', '.join(must)}." if must else "")
                                  + (" Use the bank's terms: " + "; ".join(f"{a} = {b}" for a, b in terms) + "." if terms else "") + f"\n\n{s['text']}", 1500)
                better = translating.check(s["text"], again.strip(), to)
                if again.strip() and (better["status"] == "ok" or s["check"]["status"] == "empty"):
                    s["target"], s["check"] = again.strip(), better
        if retried:
            trace.add("rules", f"{retried} segment(s) translated again after the read-back")
    except httpx.HTTPError as exc:
        raise HTTPException(503, "The model was unavailable while the text was being translated. Please try again shortly.") from exc
    for i in todo:
        if segs[i].get("target"):
            segs[i]["target"] = translating.finish(segs[i]["target"], to, opts)
    answer = translating.assemble(segs)
    rows = [{"n": n, "source": s["text"], "target": s.get("target", ""), **(s.get("check") or {"status": "ok" if s.get("copy") else "empty", "missing": [], "extra": [], "terms": []})}
            for n, s in enumerate(segs, start=1)]
    done = sum(bool(x["target"]) for x in rows)
    flagged = [x for x in rows if x["status"] == "check"]
    report = {"direction": ("ar" if to == "en" else "en") + "→" + to, "target": to, "total": len(rows), "translated": done, "to_check": len(flagged),
              "complete": done == len(rows), "left_out": len(rows) - done, "source": what, "segments": rows}
    trace.add("rules", f"Read back: {done} of {len(rows)} segment(s) translated, {len(flagged)} to check",
              [{"n": x["n"], "missing": x["missing"], "extra": x["extra"], "terms": x["terms"]} for x in flagged] or None)
    r.update(model=stats["model"], fallback=stats["model"] != r["alias"])
    meta = {"latency_ms": stats["ms"], "completion_tokens": stats["tokens"], "tokens_per_s": round(stats["tokens"] / max(stats["ms"] / 1000, 1e-3), 1)}
    if conv:
        conversation.add(conv, question, answer, {"route": r, "meta": meta, "translation": report},
                         {"question": request, "task": "translation", "lang": conversation.language(request, None)}, [f["info"] for f in files if (f.get("info") or {}).get("file_id")])
    return {"answer": answer, "route": r, "meta": meta, "trace": trace.to_list(), "conversation_id": body.conversation_id, "translation": report}


@app.get("/translation/options")
def translation_options(lang: str = "en"):
    return translating.options(lang)


# ------------------------------------------------------------------ summaries: the whole source, every point checked
SUMMARY_CHECK = os.getenv("D6_SUMMARY_CHECK", "1") == "1"
_section_slots = asyncio.Semaphore(2)


async def summary_context(files: list[dict], opts: dict, lang: str, trace: Trace) -> tuple[str, dict, int]:
    """The source of a summary as the model reads it, every passage numbered. A source that fits the prompt goes in whole.
    A longer one is read section by section first: each section is reduced to its points, and the summary is then written
    from all the points, so the end of a long document counts as much as its beginning. Returns (text, passages, sections)."""
    budget = FILE_CHARS if len(files) == 1 else FILES_CHARS
    _, everything = evidence.number(files, [10 ** 9] * len(files))
    if sum(len(p["prompt"]) for p in everything.values()) <= budget:
        extra, passages = files_context(files, trace, True)
        return extra, passages, 0
    groups = summarising.sections(list(everything.values()))
    focus = f" Note only what concerns: {opts['focus'].strip()[:300]}." if (opts.get("focus") or "").strip() else ""

    async def note(i: int, group: list[dict]) -> str:
        async with _section_slots:
            res = await llm.chat([{"role": "user", "content":
                f"This is section {i} of {len(groups)} of a longer source that will be summarised. List what a reader would act on or need to "
                "know from this section: decisions, obligations, deadlines, amounts, who must do what. At most 12 short points, each ending with "
                f"the number of the passage it comes from, like [S12]. Keep names, numbers and dates exactly as written.{focus} "
                f"Write the points in the language of the section.\n\n{summarising.section_text(group)}"}], max_tokens=900, temperature=0)
        return f"Section {i} of {len(groups)}:\n{res.text.strip()}"

    t = time.time()
    try:
        notes = await asyncio.gather(*(note(i, g) for i, g in enumerate(groups, start=1)))
    except httpx.HTTPError as exc:
        raise HTTPException(503, "The model was unavailable while the document was being read in sections. Please try again shortly.") from exc
    trace.add("files", f"Long source: read in {len(groups)} sections ({len(everything)} passages) before summarising", None, int((time.time() - t) * 1000))
    names = ", ".join(f["name"] for f in files)
    extra = (f"Source: {names}. It is too long to show whole, so it was read in {len(groups)} sections. The points noted from each "
             "section follow, each with the number of the passage it comes from.\n---\n" + "\n\n".join(notes) +
             "\n---\nWrite the summary from these points, covering every section. " + CITE)
    return extra, everything, len(groups)


async def finish_summary(answer: str, passages: dict, files: list[dict], opts: dict, lang: str, in_sections: int, trace: Trace) -> tuple[str, dict]:
    """A written summary, checked point by point against its source; points found wrong are repaired once and checked again."""
    answer = evidence.tidy(answer, passages)
    source_text = " ".join(f["text"] for f in files)
    report = {"form": opts.get("form") or "auto", "coverage": summarising.coverage(passages, files, in_sections), "legal": summarising.legal(source_text),
              "checked": False, "corrected": 0, "unsupported": [], "missing": []}
    if not SUMMARY_CHECK:
        return answer, report

    async def check(text: str) -> dict | None:
        pts = summarising.points(text)
        cited = list(dict.fromkeys(re.findall(r"S\d+", text)))
        # The whole source when it is short; otherwise the passages the summary cites, which is what its points rest on
        src = [p for p in passages.values()] if sum(len(p["prompt"]) for p in passages.values()) <= summarising.SECTION_CHARS else [passages[s] for s in cited if s in passages]
        if not pts or not src:
            return None
        try:
            data, res = await llm.chat_json([{"role": "user", "content":
                f"{summarising.CHECK}" + (f"\nThe employee asked to cover only: {opts['focus']}" if opts.get("focus") else "") +
                "\n\nSource:\n" + summarising.section_text(src) + "\n\nSummary points:\n" + "\n".join(f"{i}. {p}" for i, p in enumerate(pts, start=1))}],
                summarising.CHECK_SCHEMA, max_tokens=900)
        except (httpx.HTTPError, llm.JSONOutputError):
            return None
        bad = [{"point": pts[v["n"] - 1], "verdict": v["verdict"], "note": v.get("note", "")} for v in data.get("points", [])
               if v.get("verdict") in ("wrong", "not_in_source") and isinstance(v.get("n"), int) and 0 < v["n"] <= len(pts)]
        trace.add("llm", f"Checked {len(pts)} point(s) against the source: {len(bad)} not supported", bad or None, res.latency_ms)
        return {"bad": bad, "missing": [m for m in data.get("missing", []) if isinstance(m, str) and m.strip()][:4], "source": src}

    first = await check(answer)
    if first is None:
        return answer, report
    report.update(checked=True, unsupported=first["bad"], missing=summarising.keep_missing(first["missing"], answer))
    if first["bad"]:
        try:
            fix = await llm.chat([{"role": "user", "content":
                "A summary was checked against its source and these points were found wrong or unsupported. Rewrite the summary with those "
                "points corrected from the source, or removed if the source does not say them. Keep everything else, the form and the "
                "passage numbers in [S..] as they are. Return only the summary.\n\nProblems:\n" +
                "\n".join(f"- {b['point']} ({b['verdict']}: {b['note']})" for b in first["bad"]) +
                "\n\nSource:\n" + summarising.section_text(first["source"]) + f"\n\nSummary:\n{answer}"}], max_tokens=1500, temperature=0)
            trace.llm("Correct the unsupported points (workhorse)", fix)
            if fix.text.strip():
                answer = evidence.tidy(fix.text.strip(), passages)
                second = await check(answer)
                if second is not None:
                    report.update(corrected=len(first["bad"]) - len(second["bad"]), unsupported=second["bad"],
                                  missing=summarising.keep_missing(second["missing"], answer))
        except httpx.HTTPError:
            trace.add("rules", "Correction skipped: the model was unavailable")
    return answer, report


@app.get("/summary/options")
def summary_options(lang: str = "en"):
    return summarising.options(lang)


class ExportIn(BaseModel):
    text: str = ""
    title: str = ""
    rows: list[list[str]] = []          # [[source, translation], ...]: a two-column bilingual document


@app.post("/export.docx")
def export_text(body: ExportIn):
    """Any answer as a Word document (a summary, for one): passage marks removed, right-to-left paragraphs for Arabic."""
    return Response(to_docx(drafting.clean(body.text), body.title, body.rows), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": 'attachment; filename="summary.docx"'})


def to_docx(text: str, title: str = "", rows: list[list[str]] | None = None) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    doc = Document()

    def para(line: str, bold: bool = False) -> None:
        par = doc.add_paragraph()
        run = par.add_run(re.sub(r"\*\*|__|^#+\s*", "", line))
        run.bold = bold
        if re.search(r"[\u0600-\u06ff]", line):
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            par._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
            run._r.get_or_add_rPr().append(OxmlElement("w:rtl"))

    if title:
        para(title, bold=True)
    for line in text.split("\n") if text else []:
        para(line)
    if rows:                                # source beside translation, one row per segment
        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        for pair in rows:
            cells = table.add_row().cells
            for cell, value in zip(cells, pair[:2]):
                cell.text = re.sub(r"\*\*|__", "", value)
                if re.search(r"[\u0600-\u06ff]", value):
                    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                    cell.paragraphs[0]._p.get_or_add_pPr().append(OxmlElement("w:bidi"))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ drafting: checks, versions, approval, export
PROOFREAD = os.getenv("D6_PROOFREAD", "1") == "1"


def draft_sources(conv: dict | None, body: ChatIn | None, files: list[dict]) -> str:
    """What a draft may have taken its figures from: everything the employee wrote in this conversation, and the files."""
    msgs = conv["messages"] if conv else (body.messages if body else [])
    said = " ".join(m["content"] for m in msgs if m.get("role") == "user" and isinstance(m.get("content"), str))
    return said + " " + " ".join(f["text"] for f in files)


def draft_report(text: str, sources: str, request: str, opts: dict, files: list[dict], versions: list[dict], by: str) -> dict:
    """Everything the portal shows on a draft card: the parts, what to do before sending, where it came from."""
    report = drafting.review(text, sources, request, opts.get("type", ""))
    report["blocking"] = drafting.blocking(report)
    report.update(version=len(versions) + 1, by=by, type=opts.get("type") or "auto", approved=None)
    if versions:
        report["changes"] = drafting.diff(versions[-1]["text"], report["clean"])
    if files:
        _, numbered = evidence.number(files, [FILE_CHARS] * len(files))
        report["sources"] = evidence.attribute(report["clean"], numbered)
    return report


async def finish_draft(answer: str, question: str, r: dict, opts: dict, files: list[dict], conv: dict | None, body: ChatIn,
                       revising: bool, trace: Trace) -> tuple[str, dict]:
    """A written draft, through the policy check (one repair), the proofread and the read-back; saved as a version."""
    answer = evidence.tidy(answer, {})          # a draft carries no passage marks
    broken = [b for b in drafting.policy(answer, question, opts.get("type", "")) if not b["requested"]]
    if broken:
        trace.add("rules", f"Policy check: {len(broken)} rule(s) broken, asking for a repair", broken)
        try:        # the repair and the proofread are improvements: if the model is unavailable the draft still stands, flagged
            fix = await llm.chat([{"role": "user", "content":
                "Rewrite this document so that it no longer breaks the rules listed. Change only what is needed, keep everything "
                "else word for word, and return the whole document.\nRules broken:\n" + "\n".join(f"- {b['rule']}: \"{b['found']}\"" for b in broken) +
                f"\n\nDocument:\n{answer}"}], max_tokens=1500, temperature=0)
            trace.llm("Repair policy issues (workhorse)", fix)
            if fix.text.strip():
                answer = fix.text.strip()
        except httpx.HTTPError:
            trace.add("rules", "Policy repair skipped: the model was unavailable")
    answer, corrected = drafting.fix_weekdays(answer)
    if corrected:
        trace.add("rules", f"Calendar: {corrected} weekday(s) did not match their date and were corrected")
    open_ = drafting.placeholders(answer)
    if files and open_:
        answer = await fill_from_files(answer, open_, files, trace)
    if PROOFREAD:
        try:
            proof_ = await llm.chat([{"role": "user", "content":
                "Proofread this document. Correct spelling, grammar and punctuation mistakes only. Do not change any name, number, "
                "date, placeholder in [square brackets], meaning or wording that is already correct, and do not add or remove sentences. "
                f"Return only the corrected document.\n\n{answer}"}], max_tokens=1500, temperature=0)
            answer, changed = drafting.safe_proofread(answer, proof_.text.strip())
            trace.add("llm", f"Proofread: {changed} word(s) corrected" if changed else "Proofread: nothing changed (or the change touched a fact and was refused)",
                      proof_.meta(), proof_.latency_ms)
        except httpx.HTTPError:
            trace.add("rules", "Proofread skipped: the model was unavailable")
    versions = conv["drafts"] if conv else []
    if not revising:
        versions = []                           # a new document starts its own version history
    report = draft_report(answer, draft_sources(conv, body, files) + " " + question, question, opts, files, versions, "assistant")
    if conv:
        conv["drafts"] = versions + [{"version": report["version"], "by": "assistant", "at": time.time(), "text": report["clean"]}]
    trace.add("rules", f"Draft v{report['version']} read back: {len(report['placeholders'])} placeholder(s), {len(report['policy'])} policy issue(s), "
              f"{len(report['wrong_weekdays'])} weekday mismatch(es), {len(report['unsourced_figures'])} figure(s) with no source",
              {k: report[k] for k in ("placeholders", "policy", "wrong_weekdays", "unsourced_figures", "blocking")})
    await log_event("draft_created", r["alias"], {"conversation": body.conversation_id, "version": report["version"], "revision": revising,
                                                 "blocking": report["blocking"], "unsourced_figures": report["unsourced_figures"]})
    return report["clean"], {"draft": report}


async def fill_from_files(draft: str, open_: list[str], files: list[dict], trace: Trace) -> str:
    """Placeholders the attached files can answer (the manager's name is in the workbook) are filled in, and only with a
    value that is in the files word for word; the rest stay for the employee."""
    text = "\n\n".join(f"=== {f['name']} ===\n{f['text'][:FILE_CHARS]}" for f in files)
    schema = {"type": "object", "required": ["values"], "properties": {"values": {"type": "array", "items": {
        "type": "object", "required": ["placeholder", "value"], "properties": {"placeholder": {"type": "string"}, "value": {"type": "string"}}}}}}
    try:
        data, res = await llm.chat_json([{"role": "user", "content":
            "A draft document has gaps written as [placeholders]. For each placeholder listed, look in the files for what belongs there, "
            "given how the draft uses it. Give the value copied exactly from the files, or an empty string when the files do not state it. "
            f"Do not guess.\n\nPlaceholders: {json.dumps(open_, ensure_ascii=False)}\n\nDraft:\n{draft}\n\nFiles:\n{text}"}], schema, max_tokens=400)
    except (httpx.HTTPError, llm.JSONOutputError):
        return draft
    everything, filled = evidence.canon(text), []
    for item in data.get("values") or []:
        name, value = str(item.get("placeholder", "")).strip("[] "), str(item.get("value", "")).strip()
        if name in open_ and 1 < len(value) <= 80 and evidence.canon(value) in everything:
            draft = draft.replace(f"[{name}]", value)
            filled.append({name: value})
    trace.add("llm", f"Looked in the files for {len(open_)} placeholder(s): {len(filled)} filled in", filled, res.latency_ms)
    return draft


@app.get("/drafting/options")
def drafting_options(lang: str = "en"):
    return drafting.options(lang)


class DraftEdit(BaseModel):
    conversation_id: str
    text: str = ""
    by: str = "employee"


def _draft_message(conv: dict) -> dict:
    """The message that holds the document being drafted: the last assistant reply that is a draft."""
    for m in reversed(conv["messages"]):
        if m["role"] == "assistant" and (m.get("reply") or {}).get("draft"):
            return m
    raise HTTPException(404, {"code": "no_draft", "message": "There is no draft in this conversation."})


@app.post("/draft/save")
async def draft_save(body: DraftEdit):
    """The employee edited the draft: read it back like any other version and keep it as the newest."""
    conv = conversation.get(body.conversation_id)
    if not conv:
        raise HTTPException(404, {"code": "conversation_gone", "message": "This conversation is no longer available."})
    msg = _draft_message(conv)
    files = [FILES[f["file_id"]] for f in conv["files"] if f.get("file_id") in FILES]
    old = msg["reply"]["draft"]
    report = draft_report(body.text, draft_sources(conv, None, files), conv["last"]["question"] if conv["last"] else "", {"type": old.get("type")}, files, conv["drafts"], "employee")
    conv["drafts"].append({"version": report["version"], "by": "employee", "at": time.time(), "text": report["clean"]})
    msg["content"], msg["reply"]["draft"] = report["clean"], report
    await log_event("draft_edited", body.by, {"conversation": body.conversation_id, "version": report["version"], "blocking": report["blocking"]})
    return {"answer": report["clean"], "draft": report}


@app.post("/draft/approve")
async def draft_approve(body: DraftEdit):
    """Record that a person approved the latest version. Refused while something still has to be dealt with."""
    conv = conversation.get(body.conversation_id)
    if not conv:
        raise HTTPException(404, {"code": "conversation_gone", "message": "This conversation is no longer available."})
    draft = _draft_message(conv)["reply"]["draft"]
    if draft["blocking"]:
        raise HTTPException(409, {"code": "not_ready", "message": "This draft cannot be approved yet: " + "; ".join(draft["blocking"]) + "."})
    draft["approved"] = {"by": body.by, "at": time.strftime("%Y-%m-%d %H:%M")}
    await log_event("human_decision", body.by, {"conversation": body.conversation_id, "action": "draft_approved", "version": draft["version"],
                                               "unsourced_figures": draft["unsourced_figures"], "chars": len(draft["clean"])})
    return {"draft": draft}


@app.get("/draft/{cid}/versions")
def draft_versions(cid: str):
    conv = conversation.get(cid)
    if not conv:
        raise HTTPException(404, {"code": "conversation_gone", "message": "This conversation is no longer available."})
    return {"versions": [{**v, "at": time.strftime("%Y-%m-%d %H:%M", time.localtime(v["at"]))} for v in conv["drafts"]]}


@app.get("/draft/{cid}/export.docx")
def draft_export(cid: str, version: int = 0):
    """The draft as a Word document: the subject as a heading, right-to-left paragraphs for Arabic."""
    conv = conversation.get(cid)
    if not conv or not conv["drafts"]:
        raise HTTPException(404, {"code": "no_draft", "message": "There is no draft in this conversation."})
    v = next((d for d in conv["drafts"] if d["version"] == version), conv["drafts"][-1])
    p = drafting.parts(v["text"])
    return Response(to_docx(p["body"], p["subject"]), media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    headers={"Content-Disposition": f'attachment; filename="draft-v{v["version"]}.docx"'})


@app.get("/file/{file_id}/page/{page}")
async def file_page(file_id: str, page: int):
    """The image of one page of an uploaded PDF (or the uploaded image itself), for the source panel."""
    f = FILES.get(file_id)
    if not f or not f.get("raw"):
        raise HTTPException(404, "no page image for this file")
    cache = f.setdefault("page_png", {})
    if page not in cache:
        if f["raw_kind"] == "pdf":
            images = await platform.render_pdf(f["raw"], dpi=110, pages=[page])
            if not images:
                raise HTTPException(404, "no such page")
            cache[page] = base64.b64decode(images[0])
        else:
            buf = io.BytesIO()
            ImageOps.exif_transpose(Image.open(io.BytesIO(f["raw"]))).convert("RGB").save(buf, "PNG")
            cache[page] = buf.getvalue()
    return Response(cache[page], media_type="image/png", headers={"Cache-Control": "private, max-age=3600"})


@app.post("/route")
async def route_only(body: ChatIn):
    return await route(body.messages[-1]["content"], bool(body.file_id or body.file_ids))


_vision_slots = asyncio.Semaphore(2)     # pages are transcribed in parallel; more than two trips hosted rate limits


async def transcribe(png_base64: str) -> tuple[str, bool]:
    """One page image -> its text. The flag says the answer hit the token limit, so the page may be incomplete."""
    messages = [{"role": "user", "content": [llm.image_part(png_base64), {"type": "text", "text": TRANSCRIBE}]}]
    async with _vision_slots:
        for pause in (4, 12, 30, None):     # a busy or rate-limited model usually clears within seconds
            try:
                res = await llm.chat(messages, "vision", max_tokens=TRANSCRIBE_TOKENS, temperature=0)
                break
            except httpx.HTTPError as exc:
                if pause is None:
                    raise HTTPException(503, "The model that reads page images is busy or unavailable. Please try again shortly.") from exc
                await asyncio.sleep(pause)
    text = re.sub(r"^```\w*\n|\n```$", "", res.text.strip())      # the model sometimes fences the whole page
    return text, res.completion_tokens >= TRANSCRIBE_TOKENS


def _quiet(stage: str, done: int | None = None, total: int | None = None) -> None:
    """Progress reporter for callers that are not watching."""


async def read_image(data: bytes, report=_quiet) -> dict:
    """A photo or screenshot of a document: upright (phone photos carry their rotation in EXIF), bounded size, PNG."""
    try:
        img = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    except Exception as exc:  # noqa: BLE001 - any unreadable image is the caller's problem, not a server error
        raise intake.Rejected("damaged", kind="image") from exc
    img.thumbnail((MAX_IMAGE_SIDE, MAX_IMAGE_SIDE))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    report("reading_pages", 0, 1)
    text, cut = await transcribe(base64.b64encode(buf.getvalue()).decode())
    report("reading_pages", 1, 1)
    return {"parts": [{"page": 1, "text": text}], "paged": False, "how": "vision model (image)", "transcribed": [1],
            "cut_short": [1] if cut else []}


async def read_pdf(data: bytes, conv: dict, report=_quiet) -> dict:
    """A page with no text layer is a scan and is transcribed, page by page. Typed pages follow the converter's
    verdict on the whole document, as before: transcribed when it is mostly Arabic (`rtl_pdf`), converted otherwise.
    The result stays one entry per page, so answers can name the page."""
    pages = conv.get("page_text") or []
    converted = {p["page"]: p["markdown"].strip() for p in conv.get("page_markdown") or []}
    scanned = [p["page"] for p in pages if len(re.findall(r"\w", p["text"])) < MIN_PAGE_CHARS]
    arabic = [p["page"] for p in pages if p["page"] not in scanned] if conv.get("rtl_pdf") else []
    need = sorted(scanned + arabic)
    todo = need[:MAX_VISION_PAGES]
    images = await platform.render_pdf(data, pages=todo) if todo else []
    read_so_far = 0

    async def page(img: str) -> tuple[str, bool]:
        nonlocal read_so_far
        out = await transcribe(img)
        read_so_far += 1
        report("reading_pages", read_so_far, len(todo))
        return out

    if todo:
        report("reading_pages", 0, len(todo))
    results = dict(zip(todo, await asyncio.gather(*(page(img) for img in images))))
    # Pages past the budget: a scanned one stays unread, an Arabic one falls back to its unreliable text layer
    unread = [n for n in need[MAX_VISION_PAGES:] if n in scanned]
    parts = [{"page": p["page"], "text": results[p["page"]][0] if p["page"] in results else converted.get(p["page"]) or p["text"].strip()}
             for p in pages if p["page"] not in unread]
    if not need:
        how = "Docling"
    else:
        kinds = [f"{len(k)} {label}" for k, label in ((scanned, "scanned"), (arabic, "Arabic")) if k]
        how = ("vision model" if len(need) == len(pages) else "Docling + vision model") + \
            f" ({len(todo)} of {len(pages)} page(s) transcribed: {', '.join(kinds)})" + (f"; {len(unread)} page(s) not read" if unread else "")
    return {"parts": [p for p in parts if p["text"]], "paged": True, "how": how, "transcribed": todo, "unread": unread,
            "cut_short": [n for n, (_, cut) in results.items() if cut]}


def read_converted(conv: dict) -> dict:
    """Word, PowerPoint, HTML, Markdown, text: the converter's output. Slides count as pages; the rest have none."""
    slides = [{"page": p["page"], "text": p["markdown"].strip()} for p in conv.get("page_markdown") or [] if p["markdown"].strip()]
    if len(slides) > 1:
        return {"parts": slides, "paged": True, "how": "Docling"}
    return {"parts": [{"page": 1, "text": conv["markdown"]}], "paged": False, "how": "Docling"}


def file_text(parts: list[dict], paged: bool) -> str:
    """The text the model sees. Page (or sheet) markers let it say where an answer is; a file without pages gets none."""
    if not paged:
        return "\n\n".join(p["text"] for p in parts)
    return "\n\n".join(f"[{p.get('label') or 'Page ' + str(p['page'])}]\n{p['text']}" for p in parts)


def outline(parts: list[dict]) -> list[dict]:
    """Headings in reading order with their page: the file's table of contents, as far as the conversion kept one."""
    if parts and parts[0].get("label"):
        return [{"page": p["page"], "heading": p["label"]} for p in parts]
    return [{"page": p["page"], "heading": m.group(1).strip()} for p in parts
            for m in re.finditer(r"^#{1,6}\s+(.{1,120})$", p["text"], re.M)]       # longer than that is a paragraph, not a title


@app.get("/limits")
def upload_limits(lang: str = "en"):
    """What the portal may attach, so it can refuse an oversized or unsupported file before uploading it."""
    return intake.limits(lang)


async def receive(file: UploadFile, lang: str) -> tuple[str, bytes]:
    """The checks that need no conversion: size, type, content. Raises intake.Rejected."""
    name = file.filename or "file"
    if file.size and file.size > intake.MAX_MB * 1024 * 1024:           # known from the upload itself: refuse before reading it in
        raise intake.Rejected("too_large", lang, size=f"{file.size / 1024 / 1024:.1f}", limit=f"{intake.MAX_MB:g}")
    data = await file.read()
    intake.check(name, data, lang)
    return name, data


@app.post("/upload")
async def upload(file: UploadFile = File(...), lang: str = "en"):
    """Read a file in one blocking call. The portal uses /upload/start instead, to show progress."""
    try:
        return await read_upload(*await receive(file, lang), lang)
    except intake.Rejected as no:
        raise HTTPException(no.status, no.detail) from no


# ---- uploads as jobs: the portal starts one, then asks for its status twice a second while pages are being read
JOBS: dict[str, dict] = {}
STAGE_PERCENT = {"checking": 3, "converting": 10, "finishing": 97, "done": 100}


def job_status(job: dict) -> dict:
    done, total = job.get("pages_done"), job.get("pages_total")
    percent = 15 + round(80 * done / total) if job["stage"] == "reading_pages" and total else STAGE_PERCENT.get(job["stage"], 0)
    elapsed = (job.get("ended") or time.time()) - job["started"]
    # Seconds left: set each time a page comes back (see run_job), so it counts down in between; unknown before the first page
    eta = max(1, round(job["finish_at"] - time.time())) if job["stage"] == "reading_pages" and job.get("finish_at") else None
    return {"job_id": job["id"], "name": job["name"], "stage": job["stage"], "pages_done": done, "pages_total": total,
            "percent": percent, "elapsed_s": round(elapsed, 1), "eta_s": eta, "result": job.get("result"), "error": job.get("error")}


async def run_job(job: dict, data: bytes, lang: str) -> None:
    def report(stage: str, done: int | None = None, total: int | None = None) -> None:
        if stage == "reading_pages" and job["stage"] != "reading_pages":
            job["reading_since"] = time.time()
        job["stage"] = stage
        if total is not None:                   # later stages keep the last page count on show
            job.update(pages_done=done, pages_total=total)
            if done:                            # expected end, from the pace of the pages read so far
                job["finish_at"] = time.time() + (time.time() - job["reading_since"]) / done * (total - done)
    try:
        job["result"] = await read_upload(job["name"], data, lang, report)
        job["stage"] = "done"
    except intake.Rejected as no:
        job.update(stage="failed", error=no.detail)
    except HTTPException as exc:                # the page-reading model stayed busy through the retries
        job.update(stage="failed", error=exc.detail if isinstance(exc.detail, dict) else {"code": "busy", "message": str(exc.detail)})
    except asyncio.CancelledError:
        job["stage"] = "cancelled"
        raise
    except Exception:  # noqa: BLE001 - a job must always end in a state the portal can show
        job.update(stage="failed", error=intake.Rejected("unreadable", lang).detail)
    finally:
        job["ended"] = time.time()


@app.post("/upload/start")
async def upload_start(file: UploadFile = File(...), lang: str = "en"):
    """Check the file now (a bad one is refused in this call), then read it in the background."""
    try:
        name, data = await receive(file, lang)
    except intake.Rejected as no:
        raise HTTPException(no.status, no.detail) from no
    for old in [k for k, j in JOBS.items() if j.get("ended") and time.time() - j["ended"] > 1800]:      # finished half an hour ago
        del JOBS[old]
    job = {"id": uuid.uuid4().hex[:12], "name": name, "stage": "checking", "started": time.time()}
    JOBS[job["id"]] = job
    job["task"] = asyncio.create_task(run_job(job, data, lang))
    return job_status(job)


@app.get("/upload/{job_id}/status")
def upload_status(job_id: str):
    if job_id not in JOBS:
        raise HTTPException(404, {"code": "job_gone", "message": "This upload is no longer available."})
    return job_status(JOBS[job_id])


@app.delete("/upload/{job_id}")
async def upload_cancel(job_id: str):
    """Stop reading a file the user no longer wants; pages already sent to the model are not charged back."""
    job = JOBS.get(job_id)
    if job and not job.get("ended"):
        job["task"].cancel()
    return {"cancelled": bool(job)}


async def read_upload(name: str, data: bytes, lang: str, report=_quiet) -> dict:
    started = time.time()
    suffix = Path(name).suffix.lower()
    if suffix in IMAGE_TYPES:
        read = await read_image(data, report)
    elif suffix in sheets.TYPES:
        try:
            read = {**sheets.read(name, data), "paged": True, "kind": "sheets"}
        except sheets.SheetError as exc:
            raise intake.Rejected("no_text" if "No data" in str(exc) else "unreadable", lang) from exc
    elif suffix in (".txt", ".md"):
        # Plain text is read as it is: the converter joins the lines of a .txt file into one paragraph, which loses
        # an e-mail's header lines, a signature block, a list
        for enc in ("utf-8-sig", "cp1256", "latin-1"):
            try:
                plain = data.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        read = {"parts": [{"page": 1, "text": plain.replace("\r\n", "\n").strip()}], "paged": False, "how": "read as text"}
    else:
        report("converting")
        try:
            conv = await platform.convert_document(name, data)
        except httpx.HTTPStatusError as exc:                             # passed the checks, but the converter still could not read it
            raise intake.Rejected("unreadable", lang) from exc
        read = await read_pdf(data, conv, report) if suffix == ".pdf" else read_converted(conv)
    report("finishing")
    parts, paged = read["parts"], read["paged"]
    text = file_text(parts, paged)
    if not any(p["text"].strip() for p in parts):
        raise intake.Rejected("no_text", lang)
    heads = outline(parts)
    tables = sum(len(re.findall(r"(?:^|\n)\|[^\n]*\|\n\|[\s:|-]+\|", p["text"])) for p in parts)
    fid = uuid.uuid4().hex[:10]
    is_sheets = read.get("kind") == "sheets"
    FILES[fid] = {"name": name, "text": text, "parts": parts, "paged": paged, "outline": heads, "kind": read.get("kind"),
                  "passages": evidence.split(parts, paged)}
    if suffix == ".pdf" or suffix in IMAGE_TYPES:           # kept so the source panel can show the page a passage is on
        FILES[fid].update(raw=data, raw_kind="pdf" if suffix == ".pdf" else "image")
    if is_sheets:                                           # kept so code written for this file can be run on it (app/coding.py)
        FILES[fid].update(source=data, sheets=read["sheets"])
    structure = ({"sheets": len(parts), "rows": sum(s["rows"] for s in read["sheets"])} if is_sheets else
                 {"pages": len(parts) if paged else None, "headings": len(heads), "tables": tables})
    info = {"file_id": fid, "name": name, "chars": len(text), "pages": len(parts) if paged and not is_sheets else None,
            "converted_by": read["how"], "structure": structure, "read_s": round(time.time() - started, 1)}
    FILES[fid]["info"] = info                  # what the portal shows on the file chip; kept so a reloaded page can show it again
    return {**info, "outline": heads[:60], "pages_transcribed": read.get("transcribed", []), "pages_unread": read.get("unread", []),
            "pages_cut_short": read.get("cut_short", []), "preview": text[:600]}


@app.post("/admin/ingest")
async def ingest():
    docs = [json.loads(line) for line in (DOCS_DIR / "docs.jsonl").read_text(encoding="utf-8").splitlines() if line.strip()]
    await retrieval.ensure_collection(COLLECTION, recreate=True)
    n = await retrieval.upsert(COLLECTION, [d["text"] for d in docs], [{"name": d["name"], "library": d["library"]} for d in docs])
    return {"ingested": n}


# ------------------------------------------------------------------ OpenAI-compatible endpoint (Open WebUI, Dataiku LLM Mesh)
@app.get("/v1/models")
def models(authorization: str | None = Header(default=None)):
    require_key(authorization)
    now = int(time.time())
    return {"object": "list", "data": [{"id": m, "object": "model", "created": now, "owned_by": "nbe-gb10"}
                                       for m in ("auto", "workhorse", "hero")]}


@app.post("/v1/chat/completions")
async def completions(request: Request, authorization: str | None = Header(default=None)):
    require_key(authorization)
    body = await request.json()
    messages = body.get("messages", [])
    last = messages[-1]["content"] if messages else ""
    if isinstance(last, list):   # content parts
        last = " ".join(p.get("text", "") for p in last if p.get("type") == "text")
    if body.get("model", "auto") in ("auto", "nbe-auto"):      # "nbe-auto" kept for earlier Dataiku configs
        r = await route(last, False)
    else:
        alias = body["model"] if body["model"] in ("workhorse", "hero") else "workhorse"
        r = {"lang": detect_lang(last), "task": "explicit", "alias": alias, "reason": "model requested explicitly", "how": "client"}
    extra = await code_context(last, Trace()) if r["task"] == "code" else ""
    upstream = {**body, "model": r["alias"], "messages": build_messages(messages, r, extra)}
    if r["alias"] == "hero":
        upstream.setdefault("reasoning_effort", "low")
    headers = {"Authorization": f"Bearer {settings.llm_gateway_key}"}
    route_headers = {"X-Route-Task": r["task"], "X-Route-Model": r["alias"], "X-Route-Lang": r["lang"]}
    url = f"{settings.llm_gateway_url}/chat/completions"
    if body.get("stream"):
        async def relay():
            async with httpx.AsyncClient(timeout=600) as client:
                async with client.stream("POST", url, json=upstream, headers=headers) as resp:
                    async for chunk in resp.aiter_raw():
                        yield chunk
        return StreamingResponse(relay(), media_type="text/event-stream", headers=route_headers)
    async with httpx.AsyncClient(timeout=600) as client:
        resp = await client.post(url, json=upstream, headers=headers)
    data = resp.json()
    data["nbe_route"] = r
    from fastapi.responses import JSONResponse
    return JSONResponse(data, status_code=resp.status_code, headers=route_headers)


SAMPLES = Path(__file__).parent / "samples"       # small made-up files, so a suggested request can show a feature that needs one
# The active country's own copies (its cities, staff and currency: scripts/d6_samples.py), or the Egyptian originals
if (SAMPLES / os.getenv("COUNTRY", "EG").upper()).is_dir():
    SAMPLES = SAMPLES / os.getenv("COUNTRY", "EG").upper()


@app.post("/sample/{name}")
async def sample(name: str, lang: str = "en"):
    """Attach one of the sample files, read exactly as an upload of it would be."""
    if name not in {p.name for p in SAMPLES.iterdir()}:
        raise HTTPException(404, {"code": "no_sample", "message": "There is no such sample file."})
    try:
        return await read_upload(name, (SAMPLES / name).read_bytes(), lang)
    except intake.Rejected as no:
        raise HTTPException(no.status, no.detail) from no


@app.get("/suggested")
def suggested():
    """One request per thing the assistant does. `file` names a sample the portal attaches before it asks."""
    return [
        {"task": "file_chat", "file": "policy_ar.md", "q": "ما الحد الأقصى لقيمة العملية الواحدة، وماذا يحدث إذا زادت عليه؟"},
        {"task": "file_chat", "file": "branch_loans.xlsx", "q": "Which branch has the highest total of approved loans, and how much?"},
        {"task": "summarise", "file": "meeting_ar.txt", "q": "استخرج المهام المطلوبة من هذا المحضر في جدول: المهمة، المسؤول، الموعد."},
        {"task": "translation", "file": "thread_ar.txt", "q": "ترجم هذا الملف إلى الإنجليزية."},
        {"task": "drafting", "q": "اكتب بريداً رسمياً لمديري الفروع عن صيانة نظام الحوالات يوم الجمعة من 1 إلى 5 صباحاً"},
        {"task": "code", "file": "branch_loans.xlsx", "q": "Write pandas code that prints the total approved loan amount for each branch, highest first."},
        {"task": "code", "file": "complaints_ar.csv", "q": "اكتب كود pandas يحسب متوسط أيام المعالجة لكل نوع شكوى في هذا الملف ويطبع النتيجة."},
        {"task": "code", "q": "Explain this code and say whether it is safe:\n```python\nquery = \"SELECT * FROM customers WHERE national_id = '\" + user_input + \"'\"\ncursor.execute(query)\n```"},
    ]
