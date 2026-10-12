"""Generates the two synthetic spreadsheet test files for D6 file chat and prints their answer key.

    python make_spreadsheets.py <output folder>      (needs openpyxl)

branch_loans.xlsx   two sheets. "Loans" has a title row and a blank row above the header, as bank reports do;
                    "Branches" starts at the header.
complaints_ar.csv   Arabic headers and values, saved the way Excel saves CSV (UTF-8 with a byte-order mark).
All names and figures are invented; the random seed is fixed, so the files and answers never change.
"""
import csv
import json
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

out = Path(sys.argv[1])
rnd = random.Random(2026)

BRANCHES = [("Cairo Main", "Cairo", "Hany Mansour", 42, 60_000_000), ("Giza", "Giza", "Dina Farouk", 28, 35_000_000),
            ("Alexandria", "Alexandria", "Sherif Lotfy", 31, 40_000_000), ("Mansoura", "Mansoura", "Noha Selim", 17, 18_000_000),
            ("Aswan", "Aswan", "Tamer Abdallah", 12, 9_000_000), ("Tanta", "Tanta", "Mona Kamel", 15, 14_000_000)]
PRODUCTS = ["Personal loan", "Car loan", "Mortgage", "SME loan"]
STATUSES = ["Approved", "Approved", "Approved", "Under review", "Rejected"]
FIRST = ["Ahmed", "Mona", "Karim", "Salma", "Youssef", "Heba", "Omar", "Nour", "Hassan", "Laila"]
LAST = ["Ibrahim", "Saad", "Fathy", "Zaki", "Nabil", "Amin", "Hegazy", "Shaker"]

loans = []
for i in range(1, 41):
    product = rnd.choice(PRODUCTS)
    amount = rnd.randrange(50, 400) * 1000 if product != "Mortgage" else rnd.randrange(800, 3000) * 1000
    loans.append({"Loan ID": f"L-{1000 + i}", "Branch": rnd.choice(BRANCHES)[0], "Customer": f"{rnd.choice(FIRST)} {rnd.choice(LAST)}",
                  "Product": product, "Amount (EGP)": amount, "Rate (%)": rnd.choice([18.5, 19.25, 20.0, 21.75, 23.0]),
                  "Status": rnd.choice(STATUSES), "Application date": date(2026, 7, 1) + timedelta(days=rnd.randrange(0, 90))})

wb = Workbook()
ws = wb.active
ws.title = "Loans"
ws.append(["Retail and SME loan applications, third quarter 2026"])
ws["A1"].font = Font(bold=True, size=14)
ws.merge_cells("A1:H1")
ws.append([])
ws.append(list(loans[0]))
for row in loans:
    ws.append(list(row.values()))
wb2 = wb.create_sheet("Branches")
wb2.append(["Branch", "City", "Manager", "Staff", "Annual target (EGP)"])
for b in BRANCHES:
    wb2.append(list(b))
wb.save(out / "branch_loans.xlsx")

TYPES = ["البطاقات", "القروض", "التحويلات", "الخدمات الإلكترونية"]
STATES = ["مغلقة", "مغلقة", "قيد المعالجة", "مُصعَّدة"]
CITIES = ["القاهرة", "الجيزة", "الإسكندرية", "المنصورة"]
complaints = [{"رقم الشكوى": f"C-2026-{i:03d}", "التاريخ": (date(2026, 9, 1) + timedelta(days=rnd.randrange(0, 30))).isoformat(),
               "الفرع": rnd.choice(CITIES), "النوع": rnd.choice(TYPES), "الحالة": rnd.choice(STATES), "أيام المعالجة": rnd.randrange(1, 15)}
              for i in range(1, 31)]
with open(out / "complaints_ar.csv", "w", newline="", encoding="utf-8-sig") as f:
    w = csv.DictWriter(f, fieldnames=list(complaints[0]))
    w.writeheader()
    w.writerows(complaints)

big = max(loans, key=lambda r: r["Amount (EGP)"])
key = {
    "loan_1017_amount": next(r["Amount (EGP)"] for r in loans if r["Loan ID"] == "L-1017"),
    "loan_1017_customer": next(r["Customer"] for r in loans if r["Loan ID"] == "L-1017"),
    "loans": len(loans), "total_amount": sum(r["Amount (EGP)"] for r in loans),
    "rejected": sum(r["Status"] == "Rejected" for r in loans), "approved": sum(r["Status"] == "Approved" for r in loans),
    "largest": {"id": big["Loan ID"], "amount": big["Amount (EGP)"], "branch": big["Branch"]},
    "giza_total": sum(r["Amount (EGP)"] for r in loans if r["Branch"] == "Giza"), "giza_count": sum(r["Branch"] == "Giza" for r in loans),
    "complaints": len(complaints), "c012": next(r for r in complaints if r["رقم الشكوى"] == "C-2026-012"),
    "cards": sum(r["النوع"] == "البطاقات" for r in complaints), "escalated": sum(r["الحالة"] == "مُصعَّدة" for r in complaints),
    "avg_days": round(sum(r["أيام المعالجة"] for r in complaints) / len(complaints), 2),
    "max_days": max(r["أيام المعالجة"] for r in complaints),
}
print(json.dumps(key, ensure_ascii=False, indent=1, default=str))
